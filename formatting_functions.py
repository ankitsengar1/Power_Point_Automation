import pandas as pd
import numpy as np
import locale
from pptx.enum.dml import MSO_COLOR_TYPE,MSO_THEME_COLOR_INDEX
from pptx.enum.shapes import MSO_SHAPE_TYPE
from matplotlib.colors import LinearSegmentedColormap
from pptx.dml.color import RGBColor

#------------------------------------------------------------------------------------------------------------------------------    
def format_with_locale(value):
    try:
        locale.setlocale(locale.LC_NUMERIC, 'en_IN')
        # Round to the nearest integer and apply Indian locale formatting
        rounded_value = round(value)
        return locale.format_string("%d", rounded_value, grouping=True)
    except Exception as e:
        return value  
#---------------------------------------------------------------------------------------------------------------------------------------------

def overall_kpi(senco_ppt_xl_path,sheet_name):

    df_k = pd.read_excel(senco_ppt_xl_path,sheet_name = sheet_name)
    replacements = {'customer_count':'Total Customer Shopped', 'other_month_repeater': 'Repeaters','loyalty_sales':'Total Sales',
                    'other_month_repeater_sales':'Repeater Sales','loyalty_bills':'Total Bills','other_month_repeater_bills':'Repeater Bills',
                'loyaltyqty':'Total Qty','other_month_repeater_loyaltyqty':'Repeater Qty','visits':'Visits'}
    df_k['tag'] = df_k['tag'].replace(replacements)

    new_customers = df_k[df_k["tag"] == "current_month_onetimer"].iloc[:, 1:].values + df_k[df_k["tag"] == "current_month_repeater"].iloc[:, 1:].values
    new_customers_row = pd.DataFrame([["New Customers"] + new_customers.flatten().tolist()], columns=df_k.columns)

    new_customers_sales = df_k[df_k["tag"] == "current_month_onetimer_sales"].iloc[:, 1:].values + df_k[df_k["tag"] == "current_month_repeater_sales"].iloc[:, 1:].values
    new_customers_sales_row = pd.DataFrame([["New Customers Sales"] + new_customers_sales.flatten().tolist()], columns=df_k.columns)

    new_customers_bills = df_k[df_k["tag"] == "current_month_onetimer_bills"].iloc[:, 1:].values + df_k[df_k["tag"] == "current_month_repeater_bills"].iloc[:, 1:].values
    new_customers_bills_row = pd.DataFrame([["New Customers Bills"] + new_customers_bills.flatten().tolist()], columns=df_k.columns)

    new_customers_qty = df_k[df_k["tag"] == "current_month_onetimer_loyaltyqty"].iloc[:, 1:].values + df_k[df_k["tag"] == "current_month_repeater_loyaltyqty"].iloc[:, 1:].values
    new_customers_qty_row = pd.DataFrame([["New Customers Qty"] + new_customers_qty.flatten().tolist()], columns=df_k.columns)

    visits_per_customer = df_k[df_k["tag"] == "Visits"].iloc[:, 1:].values / df_k[df_k["tag"] == "Total Customer Shopped"].iloc[:, 1:].values
    visits_per_customer_row = pd.DataFrame([["Visits Per Customer"] + visits_per_customer.flatten().tolist()], columns=df_k.columns)

    bills_per_customer = df_k[df_k["tag"] == "Total Bills"].iloc[:, 1:].values / df_k[df_k["tag"] == "Total Customer Shopped"].iloc[:, 1:].values
    bills_per_customer_row = pd.DataFrame([["Bills Per Customer"] + bills_per_customer.flatten().tolist()], columns=df_k.columns)

    atv = df_k[df_k["tag"] == "Total Sales"].iloc[:, 1:].values / df_k[df_k["tag"] == "Total Bills"].iloc[:, 1:].values
    atv_row = pd.DataFrame([["ATV"] + atv.flatten().tolist()], columns=df_k.columns)

    abs = df_k[df_k["tag"] == "Total Qty"].iloc[:, 1:].values / df_k[df_k["tag"] == "Total Bills"].iloc[:, 1:].values
    abs_row = pd.DataFrame([["ABS"] + abs.flatten().tolist()], columns=df_k.columns)

    spend_per_visit = df_k[df_k["tag"] == "Total Sales"].iloc[:, 1:].values / df_k[df_k["tag"] == "Visits"].iloc[:, 1:].values
    spend_per_visit_row = pd.DataFrame([["Spend Per Visits"] + spend_per_visit.flatten().tolist()], columns=df_k.columns)

    asp = df_k[df_k["tag"] == "Total Sales"].iloc[:, 1:].values / df_k[df_k["tag"] == "Total Qty"].iloc[:, 1:].values
    asp_row = pd.DataFrame([["ASP"] + asp.flatten().tolist()], columns=df_k.columns)

    df_k = pd.concat([df_k, new_customers_row, new_customers_sales_row, new_customers_bills_row, new_customers_qty_row, 
                    visits_per_customer_row, bills_per_customer_row, atv_row, abs_row, spend_per_visit_row, asp_row], ignore_index=True)

    df_k = df_k[~df_k["tag"].isin(["current_month_onetimer", "current_month_repeater", "current_month_onetimer_sales", 
                                "current_month_repeater_sales","current_month_onetimer_bills","current_month_repeater_bills",
                                "current_month_onetimer_loyaltyqty","current_month_repeater_loyaltyqty",
                                "COUNT(b.uniqueitemcode)","gold_bills","diamond_bills","platinum_bills"])]

    df_k.reset_index(drop=True, inplace=True)
    df_k.rename(columns={"tag":"KPIs"},inplace = True)

    desired_order = [
        "Total Customer Shopped", "Repeaters", "New Customers", "Total Sales", "Repeater Sales",
        "New Customers Sales","Total Bills","Repeater Bills", "New Customers Bills", "Total Qty",
        "Repeater Qty","New Customers Qty", "Visits", "Visits Per Customer", "Bills Per Customer",
        "ATV", "ABS","Spend Per Visits", "ASP","gold_sales","diamond_sales","platinum_sales"
    ]

    df_k = df_k.set_index("KPIs") 
    df_k = df_k.loc[desired_order].reset_index()  
    
    return df_k
#-------------------------------------------------------------------------------------------------------------------------------------

def change_text_box(slide_1,old_text,new_text):
    for shape in slide_1.shapes:
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                if str(paragraph.text).replace("\xa0", " ").strip() == old_text:
                    first_run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
                    font = first_run.font
                    font_name = font.name
                    font_size = font.size
                    font_bold = font.bold
                    font_italic = font.italic
                    font_underline = font.underline
                    if font.color.type == MSO_COLOR_TYPE.SCHEME:
                        theme_color = font.color.theme_color
                        font_rgb_flag = 1
                    elif font.color.type == None:
                        font_rgb_flag = 0
                    elif font.color.type == MSO_COLOR_TYPE.PRESET:
                        color_index = font.color.theme_color
                        font_rgb_flag = 3
                    else:
                        rgb = font.color.rgb
                        font_rgb_flag = 2
            
                    paragraph.clear()
                    new_run = paragraph.add_run()
                    new_run.text = new_text
                    new_run.font.name = font_name
                    new_run.font.size = font_size
                    new_run.font.bold = font_bold
                    new_run.font.underline = font_underline
                    new_run.font.italic = font_italic
                    if font_rgb_flag == 1:
                        new_run.font.color.theme_color = theme_color
                    elif font_rgb_flag == 2:
                        new_run.font.color.rgb = rgb
                    elif font_rgb_flag == 3:
                        new_run.font.color.color_name = MSO_THEME_COLOR_INDEX(color_index).name

#-------------------------------------------------------------------------------------------------------------------------------------
# conditinal fomatting 
def conditional_formatting(series_list):
    
    series_list = [float(str(x)[:-1]) for x in series_list]
    sorted_numbers = sorted(series_list)

    # Define a colormap: Red -> Yellow -> Green
    colors = [
         (248/255, 105/255, 107/255),  # Red
         (254/255, 241/255, 167/255),  # Yellow
         (99/255, 190/255, 123/255)   # Green
     ]
    cmap = LinearSegmentedColormap.from_list('red_yellow_green', colors)

    # Normalize using the minimum, maximum, and center (middle) values
    min_val = min(sorted_numbers)
    max_val = max(sorted_numbers)
    mid_val = np.median(sorted_numbers)  # Middle value to anchor yellow

    # Piecewise normalization to emphasize the middle value
    def piecewise_normalize(value):
        if value <= mid_val:
            # Normalize values below the mid value (0 to 0.5)
            return 0.5 * (value - min_val) / (mid_val - min_val)
        else:
            # Normalize values above the mid value (0.5 to 1)
            return 0.5 + 0.5 * (value - mid_val) / (max_val - mid_val)

    # Map values to colors using the colormap and normalization
    colors_dict = {}
    for value in sorted_numbers:
        normalized_value = piecewise_normalize(value)
        rgba = cmap(normalized_value)  # Get RGBA color
        colors_dict[value] = tuple(int(c * 255) for c in rgba[:3])  # Convert to RGB

    return colors_dict

# def heatmap_formatting(series_list):

#     series_list = [float(str(x)[:-1]) for x in series_list]
#     sorted_numbers = sorted(series_list)  # Sort in ascending order
#     #print(sorted_numbers)

#     colors = [
#         (248/255, 105/255, 107/255),  # Red
#         (220/255, 225/255, 130/255),  # Yellow
#         (99/255, 190/255, 123/255)   # Green
#     ]

#     cmap = LinearSegmentedColormap.from_list('red_yellow_green_custom', colors)

#     # Normalize the numbers to a range between 0 and 1
#     norm = plt.Normalize(vmin=min(sorted_numbers), vmax=max(sorted_numbers))

#     # Create a colormap for each number
#     colors = cmap(norm(sorted_numbers))

#     rgb_colors = []
#     for color in colors:
#         rgb_colors.append(tuple(int(c * 255) for c in color))

#     colors_dict = {}
#     for i in range(len(sorted_numbers)):
#         colors_dict[sorted_numbers[i]]=(rgb_colors[i][0],rgb_colors[i][1],rgb_colors[i][2])

#     #print(colors_dict)
#     return colors_dict

# def green_formatting(series_list):

#     series_list = [float(str(x)[:-1]) for x in series_list]
#     sorted_numbers = sorted(series_list)

#     colors = [
#         (254/255, 241/255, 167/255),   # Yellow
#         # (254/255, 241/255, 167/255),  # Yellow
#         (99/255, 190/255, 123/255)  # Green
#     ]

#     cmap = LinearSegmentedColormap.from_list('green_yellow_white', colors)

#     # Normalize the numbers to a range between 0 and 1
#     norm = plt.Normalize(vmin=min(sorted_numbers), vmax=max(sorted_numbers))

#     # Create a colormap for each number
#     colors = cmap(norm(sorted_numbers))

#     rgb_colors = []
#     for color in colors:
#         rgb_colors.append(tuple(int(c * 255) for c in color))

#     colors_dict = {}
#     for i in range(len(sorted_numbers)):
#         colors_dict[sorted_numbers[i]]=(rgb_colors[i][0],rgb_colors[i][1],rgb_colors[i][2])

#     #print(colors_dict)
#     return colors_dict