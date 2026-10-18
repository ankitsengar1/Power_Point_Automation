
import pandas as pd
import warnings
warnings.simplefilter(action='ignore', category=FutureWarning)
from pptx.dml.color import MSO_COLOR_TYPE
from pptx.enum.dml import MSO_THEME_COLOR_INDEX
from formatting_functions import format_with_locale

def edit_table(slide_4,df_1,df_2):
    #code to check index position of table
        # table_indices = []
        # for index, shape in enumerate(slide_4.shapes):
        #     if shape.has_table:
        #         table_indices.append(index)
        # print("######################################################")
        # print(table_indices)
        # print("#####################################################")

    for i in [1,2]:
        if slide_4.shapes[i].has_table:
            table = slide_4.shapes[i].table
            for col in range(0,len(table.columns)):
                for row in range(0,len(table.rows)):
                    for paragraph in table.cell(row,col).text_frame.paragraphs:
                        first_run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
                        font = first_run.font
                        font_name = font.name
                        font_size = font.size
                        font_bold = font.bold
                        font_italic = font.italic
                        font_underline = font.underline
                        text = paragraph.text
                        if font.color.type == MSO_COLOR_TYPE.SCHEME:
                            theme_color = font.color.theme_color
                            font_rgb_flag = 1
                        elif font.color.type == None:
                            font_rgb_flag = 0
                        elif font.color.type == MSO_COLOR_TYPE.PRESET:
                            color_index = font.color.theme_color
                            font_rgb_flag = 3
                        else:
                            rgb = font.color.rgb
                            font_rgb_flag = 2

                        paragraph.clear()                    
                        new_run = paragraph.add_run()
                        new_run.font.name = font_name
                        new_run.font.size = font_size
                        new_run.font.bold = font_bold
                        new_run.font.underline = font_underline
                        new_run.font.italic = font_italic
                        # print(font_rgb_flag)
                        if font_rgb_flag == 1:
                            new_run.font.color.theme_color = theme_color
                        elif font_rgb_flag == 2:
                            new_run.font.color.rgb = rgb
                        elif font_rgb_flag == 3:
                            new_run.font.color.color_name = color_name = MSO_THEME_COLOR_INDEX(color_index).name

                        if i==1:
                            # print(row,col)
                            if row > 0 and col>=0: 
                                new_run.text = str(df_1.iloc[row-1,col])

                            elif row == 0 and col>=0:
                                #print(str(list(df_win.columns.values)[col]))
                                new_run.text = str(list(df_1.columns.values)[col])
                
                        elif i==2:
                            if row > 0 and col>=0: 
                                new_run.text =str(df_2.iloc[row-1,col])

                            elif row == 0 and col>=0:
                                #print(str(list(df_win.columns.values)[col]))
                                new_run.text = str(list(df_2.columns.values)[col])
    return slide_4


def create_df(senco_ppt_xl_path,sheet_name):
    df = pd.read_excel(senco_ppt_xl_path,sheet_name = sheet_name)
    #--------------------------------------------------------------------------------------------------------------------------------------
    # first table creation
    df_1 = df[df['type'].isin(['active', 'dormant', 'lapsed'])]
    new_index = {"active": "Active", "dormant": "Dormant", "lapsed": "Lapsed"}
    df_1.loc[:,df_1.columns[0]] = df_1.loc[:,df_1.columns[0]].replace(new_index)
    total_customers = df_1['cust_count'].sum()
    total_sales = df_1['sales'].sum()

    # Create the new columns based on requirements
    df_1.rename(columns = {'type':'Customer Type'},inplace = True)
    df_1['% Customer distribution'] = (df_1['cust_count'] / total_customers * 100).round(1).astype(str) + '%'
    df_1['% Sales distribution'] = (df_1['sales'] / total_sales * 100).round(1).astype(str) + '%'
    df_1['Customers count'] = df_1['cust_count'].apply(format_with_locale)
    df_1['ATV'] = df_1['ATV'].apply(format_with_locale)
    df_1['Available points'] = df_1['t_available_points'].apply(format_with_locale)
    df_1['Avg. Recency'] = df_1['total_avg_recency'].apply(format_with_locale)

    final_df = df_1[['Customer Type', 'Customers count', '% Customer distribution', 
                        '% Sales distribution', 'ATV', 'Available points', 'Avg. Recency']].reset_index(drop = True)
    final_df.iloc[2,5] = "-"
    #-----------------------------------------------------------------------------------------------------------------------------------------
    # second table creation
    df_2 = df[df['type'].isin(['active_onetimer', 'active_repeater','dormant_onetimer','dormant_repeater','lapsed_onetimer','lapsed_repeater'])]
    total_customers_2 = df_2['cust_count'].sum()
    total_sales_2 = df_2['sales'].sum()

    # Create the new columns based on requirements
    df_2.rename(columns = {'type':'Customer Type','total_avg_visits':'Avg. Visits'},inplace = True)
    df_2['% Customer distribution'] = (df_2['cust_count'] / total_customers_2 * 100).round(1).astype(str) + '%'
    df_2['% Sales distribution'] = (df_2['sales'] / total_sales_2 * 100).round(1).astype(str) + '%'
    df_2['Customers count'] = df_2['cust_count'].apply(format_with_locale)
    df_2['ATV'] = df_2['ATV'].apply(format_with_locale)
    df_2['Avg. Visits'] = df_2['Avg. Visits'].round(2)
    df_2['Avg. Recency'] = df_2['total_avg_recency'].apply(format_with_locale)
    final_df_2 = df_2[['Customer Type', 'Customers count', '% Customer distribution', 
                            '% Sales distribution', 'ATV','Avg. Visits', 'Avg. Recency']]
    
    return final_df,final_df_2