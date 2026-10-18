import pandas as pd
import matplotlib.pyplot as plt
from matplotlib.ticker import PercentFormatter, FuncFormatter
from matplotlib import font_manager
import matplotlib.pyplot as plt
import matplotlib.cm as cm
import numpy as np
from matplotlib.colors import LinearSegmentedColormap
# import datetime as dt
import numpy as np
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from pptx.enum.dml import MSO_COLOR_TYPE
from pptx.enum.dml import MSO_THEME_COLOR_INDEX
from pptx.util import Pt,Inches


def SubElement(parent, tagname, **kwargs):
        element = OxmlElement(tagname)
        element.attrib.update(kwargs)
        parent.append(element)
        return element

def _set_cell_border(cell,border_color, border_width='5000'):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for lines in ['a:lnL','a:lnR','a:lnT','a:lnB']:
        # if lines=='a:lnB':
        ln = SubElement(tcPr, lines, w=border_width, cap='flat', cmpd='sng', algn='ctr')
        solidFill = SubElement(ln, 'a:solidFill')
        srgbClr = SubElement(solidFill, 'a:srgbClr', val=border_color)
        prstDash = SubElement(ln, 'a:prstDash', val='solid')
        round_ = SubElement(ln, 'a:round')
        headEnd = SubElement(ln, 'a:headEnd', type='none', w='med', len='med')
        tailEnd = SubElement(ln, 'a:tailEnd', type='none', w='med', len='med')



def create_table(slide,df,df_old_colors_dict):
    df_original = df
    if len(df)>=40:
        df_1 = df[:20]
        df_2 = df[-20:]
    elif len(df)>=12:
        df_1 = df[:int(len(df)/2)]
        df_2 = df[int(len(df)/2):]
    else:
        df_1 = df

    # print(df_1)
    # print(df_2)
    df =df_1
    rows, cols = df_1.shape
    table_width = Inches(8)  # Adjust based on desired width  ####
    table_height = Inches(7)  # Adjust based on desired height
    ht = 3.9 -((len(df)*0.32)/2)
    table_1 = slide.shapes.add_table(rows + 1, cols, Inches(0.25),Inches(ht),table_width, table_height)
    # table_.table.

    header_colors = [RGBColor(51, 63, 79), RGBColor(51, 63, 79), RGBColor(51, 63, 79), RGBColor(51, 63, 79)]
    
    # Add column headers
    for col in range(cols):
        _set_cell_border(table_1.table.cell(0, col), border_color="000000", border_width="12700")
        table_1.table.cell(0, col).text_frame.text = df.columns[col]
        table_1.table.cell(0, col).text_frame.paragraphs[0].font.size = Pt(9)
        table_1.table.cell(0, col).fill.solid()
        table_1.table.cell(0, col).fill.fore_color.rgb = header_colors[col % len(header_colors)]
        table_1.table.cell(0, col).text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    table_1.table.rows[0].height = Inches(0.2)

    # Populate table data
    for row in range(1, rows + 1):
        table_1.table.rows[row].height = Inches(0.32)
        for col in range(cols):
            # table_1.table.cell(row,col).text_frame.height = Inches(0.3)
            _set_cell_border(table_1.table.cell(row, col), border_color="000000", border_width="12700")
            
            if col==0:
                table_1.table.columns[col].width = Inches(3.25)
                table_1.table.cell(row, col).text_frame.text = str(df.iloc[row - 1, col])
                table_1.table.cell(row, col).text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
            else:
                table_1.table.cell(row, col).text_frame.text = str(df.iloc[row - 1, col])
                table_1.table.columns[col].width = Inches(1)

            table_1.table.cell(row,col).text_frame.margin_bottom = Inches(0.07)
            table_1.table.cell(row,col).text_frame.margin_top = Inches(0.07)

            if len(str(df.iloc[row - 1, col]))>=35:
                table_1.table.cell(row, col).text_frame.paragraphs[0].font.size = Pt(8)
            else:
                table_1.table.cell(row, col).text_frame.paragraphs[0].font.size = Pt(9)
            
            if col==2:
                table_1.table.cell(row,col).fill.solid()
                value = int(str(df.iloc[row - 1, col])[:-1])
                table_1.table.cell(row,col).fill.fore_color.rgb = RGBColor(df_old_colors_dict[value][0],df_old_colors_dict[value][1],df_old_colors_dict[value][2])
            elif col==0:
                table_1.table.cell(row, col).fill.solid()
                table_1.table.cell(row, col).fill.fore_color.rgb = RGBColor(51, 63, 79)
                table_1.table.cell(row, col).text_frame.paragraphs[0].font.color.rgb = RGBColor(255,255,255)
            else:
                table_1.table.cell(row, col).fill.solid()
                table_1.table.cell(row, col).fill.fore_color.rgb = RGBColor(255, 255, 255)
            
            if col!=0:
                table_1.table.cell(row, col).text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
            table_1.table.cell(row, 0).text_frame.paragraphs[0].font.bold = True
            # table


    if len(df_original)>=12:
        rows, cols = df_2.shape
        df = df_2
        table_width = Inches(8)  # Adjust based on desired width
        table_height = Inches(7)  # Adjust based on desired height
        table_2 = slide.shapes.add_table(rows + 1, cols, Inches(6.75),Inches(ht),table_width, table_height)
        # table_.table.
        header_colors = [RGBColor(51, 63, 79), RGBColor(51, 63, 79), RGBColor(51, 63, 79), RGBColor(51, 63, 79)]
        
        # Add column headers
        for col in range(cols):
            _set_cell_border(table_2.table.cell(0, col), border_color="000000", border_width="12700")
            table_2.table.cell(0, col).text_frame.text = df.columns[col]
            table_2.table.cell(0, col).text_frame.paragraphs[0].font.size = Pt(9)
            table_2.table.cell(0, col).fill.solid()
            table_2.table.cell(0, col).fill.fore_color.rgb = header_colors[col % len(header_colors)]
            table_2.table.cell(0, col).text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        table_2.table.rows[0].height = Inches(0.2)

        # Populate table data
        for row in range(1, rows + 1):
            table_2.table.rows[row].height = Inches(0.32)
            for col in range(cols):
                # table_2.table.cell(row,col).text_frame.height = Inches(0.3)
                _set_cell_border(table_2.table.cell(row, col), border_color="000000", border_width="12700")
                
                if col==0:
                    table_2.table.columns[col].width = Inches(3.25)
                    table_2.table.cell(row, col).text_frame.text = str(df.iloc[row - 1, col])
                    table_2.table.cell(row, col).text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
                else:
                    table_2.table.cell(row, col).text_frame.text = str(df.iloc[row - 1, col])
                    table_2.table.columns[col].width = Inches(1)

                table_2.table.cell(row,col).text_frame.margin_bottom = Inches(0.07)
                table_2.table.cell(row,col).text_frame.margin_top = Inches(0.07)
                if len(str(df.iloc[row - 1, col]))>=35:            
                    table_2.table.cell(row, col).text_frame.paragraphs[0].font.size = Pt(8)
                else:
                    table_2.table.cell(row, col).text_frame.paragraphs[0].font.size = Pt(9)

                
                if col==2:
                    table_2.table.cell(row,col).fill.solid()
                    value = int(str(df.iloc[row - 1, col])[:-1])
                    table_2.table.cell(row,col).fill.fore_color.rgb = RGBColor(df_old_colors_dict[value][0],df_old_colors_dict[value][1],df_old_colors_dict[value][2])
                elif col==0:
                    table_2.table.cell(row, col).fill.solid()
                    table_2.table.cell(row, col).fill.fore_color.rgb = RGBColor(51, 63, 79)
                    table_2.table.cell(row, col).text_frame.paragraphs[0].font.color.rgb = RGBColor(255,255,255)
                else:
                    table_2.table.cell(row, col).fill.solid()
                    table_2.table.cell(row, col).fill.fore_color.rgb = RGBColor(255, 255, 255)
                
                if col!=0:
                    table_2.table.cell(row, col).text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
                table_2.table.cell(row, 0).text_frame.paragraphs[0].font.bold = True
                # table

    return slide


def cond_formatting(df):

    if len(df)>=40:
        df_1 = df[:20]
        df_2 = df[-20:]
    else:
        df_1 = df[:int(len(df)/2)]
        df_2 = df[int(len(df)/2):]

    max_series = df_1['% Repeat Sales']
    max_series = [float(str(x)[:-1]) for x in max_series]

    min_series = df_2['% Repeat Sales']
    min_series = [float(str(x)[:-1]) for x in min_series]

    max_colors = [
        (255/255, 255/255, 255/255),   # White
        (83/255, 176/255, 108/255)  # Green
        # (176/255, 222/255, 188/255),  # light Green
        # (254/255, 241/255, 167/255),  # Yellow
        
    ]
    min_colors = [
        (248/255, 105/255, 107/255),   # Red
        (255/255, 255/255, 255/255)   # White
    ]

    max_cmap = LinearSegmentedColormap.from_list('green_white', max_colors)
    min_cmap = LinearSegmentedColormap.from_list('red_white', min_colors)

    # Normalize the numbers to a range between 0 and 1
    max_norm = plt.Normalize(vmin=min(max_series), vmax=max(max_series))
    min_norm = plt.Normalize(vmin=min(min_series), vmax=max(min_series))

    # Create a colormap for each number
    max_colors_mapping = max_cmap(max_norm(max_series))
    min_colors_mapping = min_cmap(min_norm(min_series))

    max_rgb_colors = []
    for color in max_colors_mapping:
        max_rgb_colors.append(tuple(int(c * 255) for c in color))

    min_rgb_colors = []
    for color in min_colors_mapping:
        min_rgb_colors.append(tuple(int(c * 255) for c in color))

    colors_dict = {}
    for i in range(len(max_series)):
        colors_dict[int(max_series[i])]=(max_rgb_colors[i][0],max_rgb_colors[i][1],max_rgb_colors[i][2])
    for i in range(len(min_series)):
        colors_dict[int(min_series[i])]=(min_rgb_colors[i][0],min_rgb_colors[i][1],min_rgb_colors[i][2])

    return colors_dict



def create_df(xl_path,sheet_name):
    df = pd.read_excel(xl_path,sheet_name=sheet_name)

    df.rename(columns = {'state_tag':'Tag'},inplace = True)
    df['% Repeaters'] = ((df['repeater']/df['loyalty_customer'])*100).apply(lambda x: f"{int(round(x))}%")
    df['without % Repeat Sales'] = ((df['repater_sales']/df['loyalty_sales'])*100).apply(lambda x: int(round(x)))
    df['% Repeat Sales'] = df['without % Repeat Sales'].apply(lambda x: f"{x}%")

    #print(df)

    df['Store Name'] = df['modifiedstorecode']+' '+df['modifiedstore']
    df = df[['Store Name','% Repeaters','% Repeat Sales','without % Repeat Sales','Tag']]
    
    df_old = df[df['Tag']=='Old']
    df_new = df[df['Tag']=='New']
    df_semi_old = df[df['Tag']=='Semi_Old']

    #print(df_old)

    df_old = df_old.sort_values('without % Repeat Sales',ascending=False).drop('without % Repeat Sales',axis=1)
    df_new = df_new.sort_values('without % Repeat Sales',ascending=False).drop('without % Repeat Sales',axis=1)
    df_semi_old = df_semi_old.sort_values('without % Repeat Sales',ascending=False).drop('without % Repeat Sales',axis=1)

    df_company_old = df_old[df_old['Store Name'].str.startswith('C')]
    df_company_semi_old = df_semi_old[df_semi_old['Store Name'].str.startswith('C')]
    df_company_new = df_new[df_new['Store Name'].str.startswith('C')]

    df_franchise_old = df_old[df_old['Store Name'].str.startswith('F')]
    df_franchise_semi_old = df_semi_old[df_semi_old['Store Name'].str.startswith('F')]
    df_franchise_new = df_new[df_new['Store Name'].str.startswith('F')]

    #Removing outliers
    df_company_old = df_company_old[~((df_company_old['% Repeaters'] == '100%') & (df_company_old['% Repeat Sales'] == '100%'))]
    df_company_semi_old = df_company_semi_old[~((df_company_semi_old['% Repeaters'] == '100%') & (df_company_semi_old['% Repeat Sales'] == '100%'))]
    df_company_new = df_company_new[~((df_company_new['% Repeaters'] == '100%') & (df_company_new['% Repeat Sales'] == '100%'))]
    df_franchise_old = df_franchise_old[~((df_franchise_old['% Repeaters'] == '100%') & (df_franchise_old['% Repeat Sales'] == '100%'))]
    df_franchise_semi_old = df_franchise_semi_old[~((df_franchise_semi_old['% Repeaters'] == '100%') & (df_franchise_semi_old['% Repeat Sales'] == '100%'))]
    df_franchise_new = df_franchise_new[~((df_franchise_new['% Repeaters'] == '100%') & (df_franchise_new['% Repeat Sales'] == '100%'))]

    return df_company_old,df_company_semi_old,df_company_new,df_franchise_old,df_franchise_semi_old,df_franchise_new

